"""Real pdf2docx layout recovery pilot: actual parse(), no conversion claim by proxy."""
import modal
app=modal.App("pdf2docx-layout-pilot")
vol=modal.Volume.from_name("pdf-table-benchmark")
image=modal.Image.debian_slim(python_version="3.11").pip_install("pdf2docx==0.5.8","python-docx==1.2.0")
SRC="/benchmark/datasets/real_public_pdfs"
@app.function(image=image,cpu=2,memory=4096,timeout=650,max_containers=6,volumes={"/benchmark":vol})
def pilot(name,start):
 from pdf2docx import Converter
 import json,time,traceback
 from pathlib import Path
 src=Path(SRC)/name
 target=Path("/benchmark/runs/native-layout-40-20261021")/"pdf2docx";target.mkdir(parents=True,exist_ok=True)
 t=time.perf_counter();entry={"file":name,"start":start}
 cv=None
 try:
  cv=Converter(str(src))
  # convert selected 2-page window, extract docx text and tables from native layout
  dest=target/(src.stem+"_"+str(start)+".docx")
  cv.convert(str(dest),start=start,end=start+2,multi_processing=False)
  from docx import Document
  doc=Document(str(dest))
  entry.update(ok=True,seconds=round(time.perf_counter()-t,2),paragraphs=len(doc.paragraphs),tables=len(doc.tables),table_shapes=[[len(t.rows),len(t.columns)] for t in doc.tables],paragraph_preview=[p.text[:150] for p in doc.paragraphs[:25]])
 except Exception as e:entry.update(ok=False,error=repr(e),trace=traceback.format_exc()[-700:])
 finally:
  if cv is not None:cv.close()
 (target/(src.stem+"_"+str(start)+".json")).write_text(json.dumps(entry,ensure_ascii=False))
 vol.commit()
 return entry
@app.local_entrypoint()
def main():
 tests=[("epa_wil_female_summary.pdf",4),("fda_nda_210598.pdf",5),("ntp_tr600.pdf",5),("who_annex6_legacy.pdf",7)]
 for x in pilot.starmap(tests,order_outputs=False,return_exceptions=True):print("PILOT",x,flush=True)
